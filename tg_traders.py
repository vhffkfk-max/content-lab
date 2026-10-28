"""트레이더스 방(내가 운영하는 포럼형 그룹) 게이미피케이션 '왕' 후보 집계 + 토픽(방)별 활성화 피드백.

    python tg_traders.py                                  # 10월 1일~오늘 중간 체크 (엑셀 + 사이트 통계 + Claude 피드백)
    python tg_traders.py --end 2026-10-31                 # 최종
    python tg_traders.py --no-ai                          # Claude 없이 통계만
    python tg_traders.py --from-cache                     # 텔레그램 다시 안 읽고 지난 집계로 피드백만 다시
왕 4개 (공지 기준):
  💬 활동왕 — 채팅·참여량: 메시지(10자 이상) 1 + 참여한 날 3 + 받은 반응 1 + 받은 답글 2 (tg_participation.py 와 같은 식)
  😇 성장왕 — 한 달 꾸준한 매매일지: 일지 토픽(수익 인증방)에 글·사진을 올린 '서로 다른 날' 수 → 담당자가 내용 확인
  🧠 정보왕 — 양질의 정보: 정보 토픽의 링크·전달(퍼오기)·100자 이상 글·설명 붙은 사진 → 글 2 + 받은 반응 1 + 받은 답글 1
  🤝 기여왕 — 질문 답변·회원 도움: 다른 회원 질문(?)에 단 답글 3 + 다른 회원 글에 단 답글(10자 이상) 1
관리자·봇은 순위에서 뺀다. 회원 이름이 들어간 엑셀은 output/participation/ (내 PC에만, 저장소에 안 올림).
사이트(output/site/trd_rooms.json)에는 이름·원문 없이 토픽 통계 + 피드백 + 부문별 후보 수만 올린다.
"""
import argparse
import json
import os
import re
import statistics
import sys
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from typing import List

from common import OUTPUT_DIR, load_env

CHAT_ID = -1004375748687
MODEL = "claude-opus-5-5"
KST = timezone(timedelta(hours=9))
OUT_DIR = OUTPUT_DIR / "participation"
CACHE = OUT_DIR / "traders_stats.json"          # 이름 포함 → 로컬 전용 폴더
SITE = OUTPUT_DIR / "site" / "trd_rooms.json"
MIN_CHARS = 10
JOURNAL_TOPICS = {"수익 인증방"}                  # 매매일지 올리는 곳
NOT_INFO_TOPICS = {"수익 인증방", "공지사항"}      # 정보왕 집계에서 뺄 곳
URL_RE = re.compile(r"https?://|t\.me/|www\.", re.I)
Q_RE = re.compile(r"\?|？|(나요|까요|ㄹ까|을까|인가요|있나요|되나요|어떤가요|맞나요|궁금)")


def name_of(u):
    if u is None:
        return "(알 수 없음)"
    full = " ".join(x for x in (getattr(u, "first_name", None), getattr(u, "last_name", None)) if x)
    return full or getattr(u, "title", None) or str(u.id)


def topic_of(m):
    r = m.reply_to
    if r and getattr(r, "forum_topic", False):
        return r.reply_to_top_id or r.reply_to_msg_id
    return 1


def collect(start, end):
    from telethon.tl.functions.messages import GetForumTopicsRequest
    from telethon.tl.types import ChannelParticipantsAdmins
    from tg_login import client
    c = client()
    c.connect()
    if not c.is_user_authorized():
        sys.exit("❌ 텔레그램 로그인이 풀렸어요 — python tg_login.py 로 다시 로그인해 주세요.")
    chat = c.get_entity(CHAT_ID)
    topics = {t.id: {"title": t.title, "closed": bool(getattr(t, "closed", False))}
              for t in c(GetForumTopicsRequest(peer=chat, offset_date=None, offset_id=0, offset_topic=0, limit=100)).topics}
    admins = {u.id for u in c.get_participants(chat, filter=ChannelParticipantsAdmins)}
    members = getattr(c.get_entity(chat), "participants_count", None)
    msgs, users = [], {}
    for m in c.iter_messages(chat, offset_date=end):
        if m.date < start:
            break
        if m.action or not m.sender_id or m.id in topics:
            continue
        if m.sender_id not in users:
            u = m.sender
            users[m.sender_id] = {"name": name_of(u), "username": getattr(u, "username", None) or "",
                                  "bot": bool(getattr(u, "bot", False))}
        text = m.message or ""
        r = m.reply_to
        reply_to = None
        if r and r.reply_to_msg_id and not (getattr(r, "forum_topic", False) and not r.reply_to_top_id):
            reply_to = r.reply_to_msg_id      # 토픽 첫 글(토픽 자체)에 단 건 답글이 아님
        has_url = bool(URL_RE.search(text)) or any(type(e).__name__ in ("MessageEntityUrl", "MessageEntityTextUrl")
                                                     for e in (m.entities or []))
        msgs.append({"id": m.id, "uid": m.sender_id, "topic": topic_of(m), "at": m.date.astimezone(KST).isoformat(),
                     "len": len(text.strip()), "q": bool(Q_RE.search(text)), "url": has_url, "fwd": bool(m.fwd_from),
                     "media": bool(m.media), "react": sum(x.count for x in m.reactions.results) if m.reactions else 0,
                     "reply_to": reply_to})
    c.disconnect()
    return {"title": getattr(chat, "title", "트레이더스"), "members": members, "topics": {str(k): v for k, v in topics.items()},
            "admins": sorted(admins), "users": {str(k): v for k, v in users.items()}, "msgs": msgs}


def analyze(raw, start, end):
    topics = {int(k): v for k, v in raw["topics"].items()}
    admins, users = set(raw["admins"]), {int(k): v for k, v in raw["users"].items()}
    msgs = raw["msgs"]
    by_id = {m["id"]: m for m in msgs}
    skip = lambda uid: uid in admins or users.get(uid, {}).get("bot")
    tname = lambda tid: topics.get(tid, {}).get("title", f"토픽 {tid}")

    # 답글 관계
    got_replies = Counter()      # 다른 사람에게 받은 답글
    answered = set()             # 다른 사람이 답한 글
    for m in msgs:
        t = by_id.get(m["reply_to"])
        if t and t["uid"] != m["uid"]:
            got_replies[t["id"]] += 1
            answered.add(t["id"])

    P = defaultdict(lambda: {"msgs": 0, "short": 0, "days": set(), "react": 0, "replies": 0,
                             "journal_days": set(), "journal_posts": 0,
                             "info_posts": 0, "info_react": 0, "info_replies": 0, "info_topics": Counter(),
                             "answers": 0, "helps": 0})
    for m in msgs:
        uid = m["uid"]
        p = P[uid]
        day = m["at"][:10]
        title = tname(m["topic"])
        p["days"].add(day)
        p["msgs" if m["len"] >= MIN_CHARS else "short"] += 1
        p["react"] += m["react"]
        p["replies"] += got_replies[m["id"]]
        if title in JOURNAL_TOPICS and (m["len"] >= MIN_CHARS or m["media"]):
            p["journal_days"].add(day)
            p["journal_posts"] += 1
        if title not in NOT_INFO_TOPICS and (m["url"] or m["fwd"] or m["len"] >= 100 or (m["media"] and m["len"] >= 30)):
            p["info_posts"] += 1
            p["info_react"] += m["react"]
            p["info_replies"] += got_replies[m["id"]]
            p["info_topics"][title] += 1
        t = by_id.get(m["reply_to"])
        if t and t["uid"] != uid:
            if t["q"]:
                p["answers"] += 1
            elif m["len"] >= MIN_CHARS:
                p["helps"] += 1

    elapsed = max(1, -(-int((min(end, datetime.now(KST)) - start).total_seconds()) // 86400))
    people = []
    for uid, p in P.items():
        if skip(uid):
            continue
        u = users.get(uid, {})
        people.append({
            "이름": u.get("name", str(uid)), "아이디": f"@{u['username']}" if u.get("username") else "", "회원번호": uid,
            "활동왕": p["msgs"] + 3 * len(p["days"]) + p["react"] + 2 * p["replies"],
            "참여한 날": len(p["days"]), "메시지(10자 이상)": p["msgs"], "짧은 메시지": p["short"], "받은 반응": p["react"], "받은 답글": p["replies"],
            "성장왕": len(p["journal_days"]), "일지 올린 날": len(p["journal_days"]), "일지 글": p["journal_posts"],
            "일지 꾸준함": f"{len(p['journal_days'])}/{elapsed}일",
            "정보왕": 2 * p["info_posts"] + p["info_react"] + p["info_replies"], "정보 글": p["info_posts"],
            "정보 글 반응": p["info_react"], "정보 글 답글": p["info_replies"],
            "정보 주 토픽": ", ".join(t for t, _ in p["info_topics"].most_common(2)),
            "기여왕": 3 * p["answers"] + p["helps"], "질문에 단 답": p["answers"], "도움 답글": p["helps"],
        })

    # 토픽별 지표 (이름 없음)
    last7 = (min(end, datetime.now(KST)) - timedelta(days=7)).isoformat()
    prev7 = (min(end, datetime.now(KST)) - timedelta(days=14)).isoformat()
    T = defaultdict(lambda: {"msgs": 0, "users": set(), "hours": Counter(), "last7": 0, "prev7": 0, "q": 0, "q_open": 0,
                             "admin": 0, "react": [], "media": 0, "urls": 0, "long": 0, "days": set(), "replies": 0})
    for m in msgs:
        t = T[m["topic"]]
        t["msgs"] += 1
        t["users"].add(m["uid"])
        t["days"].add(m["at"][:10])
        t["hours"][int(m["at"][11:13])] += 1
        t["react"].append(m["react"])
        t["media"] += m["media"]
        t["urls"] += m["url"] or m["fwd"]
        t["long"] += m["len"] >= 100
        t["replies"] += bool(m["reply_to"])
        t["admin"] += m["uid"] in admins
        if m["at"] >= last7:
            t["last7"] += 1
        elif m["at"] >= prev7:
            t["prev7"] += 1
        if m["q"] and not skip(m["uid"]):
            t["q"] += 1
            t["q_open"] += m["id"] not in answered
    rooms = []
    for tid, info in sorted(topics.items(), key=lambda kv: -T[kv[0]]["msgs"]):
        t = T[tid]
        n = t["msgs"]
        rooms.append({
            "id": tid, "title": info["title"], "closed": info["closed"], "msgs": n, "active_members": len(t["users"] - admins),
            "per_day": round(n / elapsed, 1), "active_days": len(t["days"]), "last7": t["last7"], "prev7": t["prev7"],
            "trend": round((t["last7"] - t["prev7"]) / t["prev7"] * 100) if t["prev7"] else None,
            "peak_hours": [h for h, _ in t["hours"].most_common(3)], "questions": t["q"], "unanswered": t["q_open"],
            "admin_share": round(t["admin"] / n * 100) if n else 0, "react_avg": round(statistics.mean(t["react"]), 1) if n else 0,
            "media_share": round(t["media"] / n * 100) if n else 0, "link_share": round(t["urls"] / n * 100) if n else 0,
            "long_share": round(t["long"] / n * 100) if n else 0, "reply_share": round(t["replies"] / n * 100) if n else 0,
            "first_day": min(t["days"]) if t["days"] else None, "last_day": max(t["days"]) if t["days"] else None,
        })
    total_members = len({m["uid"] for m in msgs} - admins)
    return people, rooms, {"elapsed_days": elapsed, "messages": len(msgs), "active_members": total_members,
                           "members": raw.get("members")}


AWARDS = [("활동왕", "💬", ["참여한 날", "메시지(10자 이상)", "받은 반응", "받은 답글"]),
          ("성장왕", "😇", ["일지 꾸준함", "일지 글"]),
          ("정보왕", "🧠", ["정보 글", "정보 글 반응", "정보 글 답글", "정보 주 토픽"]),
          ("기여왕", "🤝", ["질문에 단 답", "도움 답글"])]


def save_excel(people, title, start_s, end_s, totals):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    tops = {}
    for award, icon, cols in AWARDS:
        ranked = sorted([p for p in people if p[award] > 0], key=lambda p: (-p[award], -p["참여한 날"]))
        tops[award] = ranked
        ws = wb.create_sheet(f"{icon}{award}")
        head = ["순위", "이름", "아이디", "점수"] + cols + ["회원번호"]
        ws.append(head)
        for i, p in enumerate(ranked, 1):
            ws.append([i, p["이름"], p["아이디"], p[award]] + [p[c] for c in cols] + [p["회원번호"]])
    ws = wb.create_sheet("기준")
    for line in [("방", title), ("기간", f"{start_s} ~ {end_s} (한국시간, 중간 체크면 오늘까지)"),
                 ("뽑은 시각", datetime.now(KST).strftime("%Y-%m-%d %H:%M")),
                 ("읽은 메시지", totals["messages"]), ("활동 회원", totals["active_members"]),
                 ("💬 활동왕", "메시지(10자 이상) 1 + 참여한 날 3 + 받은 반응 1 + 받은 답글 2"),
                 ("😇 성장왕", f"{', '.join(JOURNAL_TOPICS)}에 글(10자 이상)·사진을 올린 서로 다른 날 수 — 일지 내용은 담당자 확인"),
                 ("🧠 정보왕", "정보 토픽(수익 인증방·공지 제외)의 링크·전달·100자 이상·설명 붙은 사진 글: 글 2 + 받은 반응 1 + 받은 답글 1"),
                 ("🤝 기여왕", "다른 회원 질문(? 등)에 단 답글 3 + 다른 회원 글에 단 답글(10자 이상) 1"),
                 ("제외", "관리자·봇"), ("주의", "텔레그램 '답장' 기능을 쓴 것만 답글로 셈 — 답장 없이 이어 쓴 답은 못 셈")]:
        ws.append(list(line))
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="F3D6E2")
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = max(10, min(40, max(len(str(x.value or "")) for x in col) * 1.6))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / f"왕선정_{title}_{start_s}_{end_s}.xlsx".replace("/", "_")
    wb.save(out)
    return out, tops


def feedback(rooms, totals, awards_summary, period):
    if not os.environ.get("ANTHROPIC_API_KEY"):
        print("ANTHROPIC_API_KEY 가 없어 방별 피드백을 건너뛰어요.")
        return None
    import anthropic
    from pydantic import BaseModel
    from feedback import feedback_prompt

    class Room(BaseModel):
        title: str
        status: str          # 활발 / 보통 / 조용 / 관리자 전용
        problem: str
        actions: List[str]
        post_example: str

    class Plan(BaseModel):
        summary: List[str]
        rooms: List[Room]
        award_tips: List[str]

    system = ("당신은 유료·무료 투자 커뮤니티(텔레그램 포럼형 그룹 '트레이더스') 운영 컨설턴트입니다. 운영자는 회원 참여 보상으로 "
              "매달 활동왕·성장왕(매매일지)·정보왕·기여왕을 뽑아 15만원 상당 상품을 줍니다. 토픽(방)별 통계를 보고 각 방을 살릴 "
              "구체적인 운영 피드백을 한국어로 쓰세요. 규칙: 회원 이름·아이디·원문을 지어내거나 쓰지 말 것, 투자 권유·수익 보장 문구 금지, "
              "숫자 근거를 들 것, actions 는 이번 주에 운영자가 바로 할 일 2~4개(시간대·형식까지), post_example 은 그 방에 운영자가 "
              "그대로 올릴 수 있는 글 한 편(3~6줄, 참여를 부르는 질문·미션 포함). closed=true 인 방은 관리자만 글을 쓰는 방이니 "
              "'회원 반응(반응·조회 유도)' 관점으로 조언하세요. summary 는 그룹 전체 진단 3~5줄, award_tips 는 남은 기간 "
              "왕 선정 경쟁을 띄울 운영 팁 2~4개.")
    prompt = (f"기간: {period} (경과 {totals['elapsed_days']}일) · 전체 메시지 {totals['messages']} · 활동 회원 {totals['active_members']}"
              f"명 · 그룹 인원 {totals.get('members') or '?'}명\n부문별 후보 수(점수>0): {json.dumps(awards_summary, ensure_ascii=False)}\n\n"
              "토픽별 지표 (per_day=하루 평균 글, trend=최근7일 vs 그 전 7일 %, peak_hours=글 많은 시각(한국시간), unanswered=답장 못 받은 "
              "회원 질문 수, *_share=%, first_day/last_day=기간 안 첫·마지막 글 — 10/23 무렵 토픽 개편으로 '공지사항'(예전 기본 채팅방)이 닫히고 "
              "주제별 방이 새로 열렸을 수 있음):\n" + json.dumps(rooms, ensure_ascii=False, indent=1)
              + feedback_prompt(["텔레그램 문구", "트레이더스"]))
    client = anthropic.Anthropic()
    try:
        r = client.messages.parse(model=MODEL, max_tokens=12000, system=system,
                                  messages=[{"role": "user", "content": prompt}], output_format=Plan)
    except (anthropic.APIStatusError, anthropic.APIConnectionError) as e:
        print(f"⚠️ Claude API 오류: {e} — 피드백 없이 통계만 저장해요.")
        return None
    if r.parsed_output is None:
        print(f"⚠️ 피드백을 받지 못했어요(stop_reason={r.stop_reason}).")
        return None
    print(f"✅ 방별 피드백 (입력 {r.usage.input_tokens:,} / 출력 {r.usage.output_tokens:,} 토큰)")
    return r.parsed_output.model_dump()


def stats_key(rooms):
    import hashlib
    return hashlib.sha1(json.dumps(rooms, sort_keys=True).encode()).hexdigest()[:12]


def ai_only():
    if not SITE.exists():
        return print("trd_rooms.json 이 없어요.")
    d = json.loads(SITE.read_text(encoding="utf-8"))
    if d.get("ai") and d.get("ai_key") == stats_key(d["rooms"]):
        return print("트레이더스 통계가 그대로라 피드백을 다시 쓰지 않아요.")
    ai = feedback(d["rooms"], d["totals"], {a["name"]: a["candidates"] for a in d["awards"]}, d["period"])
    if ai:
        d.update(ai=ai, ai_key=stats_key(d["rooms"]), ai_at=datetime.now(KST).isoformat(timespec="minutes"))
        SITE.write_text(json.dumps(d, ensure_ascii=False, indent=1), encoding="utf-8")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--start", default="2026-10-01")
    ap.add_argument("--end", default=None, help="끝일 YYYY-MM-DD (포함). 없으면 오늘까지 = 중간 체크")
    ap.add_argument("--no-ai", action="store_true")
    ap.add_argument("--from-cache", action="store_true")
    ap.add_argument("--ai-only", action="store_true", help="깃허브 실행용: 올라온 trd_rooms.json 통계로 피드백만 (통계가 바뀌었을 때만)")
    args = ap.parse_args()
    load_env()
    if args.ai_only:
        return ai_only()
    now = datetime.now(KST)
    start = datetime.fromisoformat(args.start).replace(tzinfo=KST)
    end_s = args.end or now.strftime("%Y-%m-%d")
    end = datetime.fromisoformat(end_s).replace(tzinfo=KST) + timedelta(days=1)
    final = bool(args.end) and end <= now
    if args.from_cache and CACHE.exists():
        raw = json.loads(CACHE.read_text(encoding="utf-8"))
    else:
        print("트레이더스 메시지 읽는 중... (몇 분 걸려요)")
        raw = collect(start, end)
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        CACHE.write_text(json.dumps(raw, ensure_ascii=False), encoding="utf-8")
    people, rooms, totals = analyze(raw, start, end)
    out, tops = save_excel(people, raw["title"], args.start, end_s, totals)
    print(f"✅ {raw['title']} {args.start}~{end_s}: 메시지 {totals['messages']}개, 활동 회원 {totals['active_members']}명 → {out}")
    for award, icon, cols in AWARDS:
        print(f"\n{icon} {award} (후보 {len(tops[award])}명)")
        for i, p in enumerate(tops[award][:5], 1):
            print(f"  {i}. {p['이름']} {p['아이디']} — {p[award]}점 ({', '.join(f'{c} {p[c]}' for c in cols)})")
    awards_summary = {a: len(tops[a]) for a, _, _ in AWARDS}
    period = f"{args.start} ~ {end_s}" + ("" if final else " (중간 체크)")
    ai = None if args.no_ai else feedback(rooms, totals, awards_summary, period)
    old = json.loads(SITE.read_text(encoding="utf-8")) if SITE.exists() else {}
    SITE.parent.mkdir(parents=True, exist_ok=True)
    SITE.write_text(json.dumps({
        "period": period, "final": final, "updated_at": now.isoformat(timespec="minutes"), "totals": totals,
        "awards": [{"name": a, "icon": i, "candidates": len(tops[a]),
                    "rule": {"활동왕": "채팅·참여량 (10자 이상 글 + 참여한 날 + 받은 반응·답글)",
                             "성장왕": "매매일지 꾸준함 (수익 인증방에 올린 날 수, 담당자 확인)",
                             "정보왕": "양질의 정보 (링크·퍼오기·긴 글·설명 붙은 사진 + 받은 반응)",
                             "기여왕": "질문 답변·회원 도움 (다른 회원 글에 단 답장)"}[a],
                    "top_score": tops[a][0][a] if tops[a] else 0} for a, i, _ in AWARDS],
        "rooms": rooms, "ai": ai or old.get("ai"), "ai_key": stats_key(rooms) if ai else old.get("ai_key"),
    }, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"✅ 사이트용 {SITE.name} (이름 없이 통계·피드백만)")


if __name__ == "__main__":
    sys.exit(main())
