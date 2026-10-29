"""카드뉴스 HTML → 캔바에서 고칠 수 있는 사본 (배경 그림 1장 + 글상자들) PPTX → 캔바 '발행 대기' 폴더.

    from publish.canva_copy import build_pptx, to_canva
    build_pptx(slides_html, css, "out.pptx")      # forms.slides_of(post) 의 HTML 그대로
    design = to_canva("out.pptx", "제목")          # → {"id", "urls": {"edit_url", ...}}

어떻게 (2026-10-29 시험):
  - 배경·차트·로고·버튼 모양은 글자만 투명하게 한 스크린샷 한 장 → 원본과 똑같이 보임 (캔바는 그림으로만 고침)
  - 글자는 크롬에서 잰 위치·크기·굵기·색 그대로 글상자 하나씩 → 캔바에서 글자·색을 바로 고침
  - PDF 로 보내 보니 캔바가 가까운 줄들을 멋대로 한 글상자로 묶어 줄이 사라지고 위치가 밀려서 PPTX 로 함
  - 그라데이션 글씨는 단색 주황, 글꼴은 캔바에 있는 Noto Sans KR / Noto Serif KR
발행 이미지는 여전히 render.render() 결과를 쓴다 (캔바 사본은 사람이 다듬는 용도).
"""
import html as htmllib
import json
import pathlib
import re
import tempfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from publish.render import H, W, chrome, page

PX = 9525
GRAD = "#ffb02e"   # 그라데이션 글씨 대신 쓸 색
SLACK = 0.06       # 캔바 글꼴 폭이 조금 달라도 줄이 안 넘어가게 글상자를 조금 넓힘

NO_TEXT = """
*{color:transparent!important;-webkit-text-fill-color:transparent!important;text-shadow:none!important;text-decoration-color:transparent!important}
[style*="background-clip:text"]{background:none!important}
"""

# 글자가 들어 있는 가장 안쪽 덩어리(블록)마다 위치·글꼴·글자 조각(색·굵기별)을 잰다
MEASURE = r"""<script>
document.fonts.ready.then(() => {
  const INL = new Set(['B','STRONG','SPAN','I','EM','U','BR','SMALL','SUP','SUB','A','MARK']);
  const out = [];
  const color = el => {
    for (let e = el; e && e.nodeType === 1; e = e.parentElement)
      if ((e.getAttribute('style') || '').includes('background-clip:text')) return '__GRAD__';
    return getComputedStyle(el).color;
  };
  const block = el => {
    let has = false;
    for (const n of el.childNodes) {
      if (n.nodeType === 3 && n.textContent.trim()) has = true;
      else if (n.nodeType === 1 && n.textContent.trim()) { if (!INL.has(n.tagName)) return false; has = true; }
    }
    return has;
  };
  // 크롬이 줄을 바꾼 자리에 그대로 줄바꿈을 넣는다 (캔바는 한글 낱말 중간에서 줄을 바꿔 버려서)
  const runs = (el, acc, st = {top: null}) => {
    for (const n of el.childNodes) {
      if (n.nodeType === 3) {
        const c = color(el), w = +getComputedStyle(el).fontWeight, s = n.textContent;
        let piece = '';
        for (const m of s.matchAll(/\S+|\s+/g)) {
          if (/^\s/.test(m[0])) { piece += ' '; continue; }
          const rg = document.createRange(); rg.setStart(n, m.index); rg.setEnd(n, m.index + m[0].length);
          const top = rg.getClientRects()[0]?.top;
          if (top != null && st.top != null && top > st.top + 4) {
            if (piece.trim()) acc.push({t: piece.trimEnd(), c, w});
            acc.push({t: '\n'}); piece = '';
          }
          if (top != null) st.top = top;
          piece += m[0];
        }
        if (piece) acc.push({t: piece, c, w});
      }
      else if (n.tagName === 'BR') { acc.push({t: '\n'}); st.top = null; }
      else if (n.nodeType === 1) runs(n, acc, st);
    }
    return acc;
  };
  for (const el of document.querySelectorAll('body *')) {
    if (el.closest('svg') || !block(el)) continue;
    if (INL.has(el.tagName) && el.parentElement && block(el.parentElement)) continue;   // 문장 안 강조(<b>)는 문장 글상자에 이미 있음
    // 글자가 실제로 차지한 곳 (그림·아이콘 빼고, 가운데 정렬 flex 도 맞게)
    let r = null;
    const tw = document.createTreeWalker(el, NodeFilter.SHOW_TEXT);
    for (let n; (n = tw.nextNode());) {
      if (!n.textContent.trim()) continue;
      const rg = document.createRange(); rg.selectNodeContents(n);
      const q = rg.getBoundingClientRect();
      r = r ? {left: Math.min(r.left, q.left), top: Math.min(r.top, q.top), right: Math.max(r.right, q.right), bottom: Math.max(r.bottom, q.bottom)}
            : {left: q.left, top: q.top, right: q.right, bottom: q.bottom};
    }
    if (!r) continue;
    r.width = r.right - r.left; r.height = r.bottom - r.top;
    const cs = getComputedStyle(el);
    if (cs.visibility === 'hidden' || +cs.opacity === 0 || r.width < 1) continue;
    const fs = parseFloat(cs.fontSize);
    out.push({x: r.left, y: r.top, w: r.width, h: r.height, fs, lh: cs.lineHeight === 'normal' ? fs * 1.3 : parseFloat(cs.lineHeight),
              ls: cs.letterSpacing === 'normal' ? 0 : parseFloat(cs.letterSpacing), align: cs.textAlign,
              serif: /HSerif|Serif KR|Serif CJK/.test(cs.fontFamily), runs: runs(el, [])});
  }
  const pre = document.createElement('pre'); pre.id = '__m'; pre.textContent = JSON.stringify(out); document.body.appendChild(pre);
});
</script>"""


def measure(body, css, tmp):
    dom = chrome(page(body + MEASURE, css), tmp, "--dump-dom")
    m = re.search(r'<pre id="__m">(.*?)</pre>', dom, re.S)
    return json.loads(htmllib.unescape(m.group(1))) if m else []


def rgb(c):
    if c == "__GRAD__":
        return RGBColor.from_string(GRAD[1:])
    v = [int(float(x)) for x in re.findall(r"[\d.]+", c)[:3]]
    return RGBColor(*v) if len(v) == 3 else RGBColor(255, 255, 255)


def px(v):
    return Emu(int(v * PX))


def add_block(slide, b):
    slack = b["w"] * SLACK
    x, w = b["x"], b["w"] + slack
    if b["align"] == "center":
        x -= slack / 2
    elif b["align"] in ("right", "end"):
        x -= slack
    tb = slide.shapes.add_textbox(px(x), px(b["y"]), px(w), px(b["h"]))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    lines, cur = [], []
    for r in b["runs"]:
        if r["t"] == "\n":
            lines.append(cur)
            cur = []
        else:
            cur.append(r)
    lines.append(cur)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT}.get(b["align"], PP_ALIGN.LEFT)
        p.line_spacing = round(b["lh"] / b["fs"], 2)
        for j, r in enumerate(line):
            t = r["t"].lstrip() if j == 0 else r["t"]
            t = t.rstrip() if j == len(line) - 1 else t
            if not t:
                continue
            run = p.add_run()
            run.text = t
            f = run.font
            name = "Noto Serif KR" if b["serif"] else "Noto Sans KR"
            f.name = name
            rpr = run._r.get_or_add_rPr()
            for tag in ("a:ea", "a:cs"):       # 한글·기호도 같은 글꼴로 (안 넣으면 캔바가 숫자·영문을 다른 글꼴로 바꿈)
                el = rpr.find(qn(tag))
                if el is None:
                    el = rpr.makeelement(qn(tag), {})
                    rpr.append(el)
                el.set("typeface", name)
            f.size = Pt(round(b["fs"] * 0.75, 1))
            f.bold = r["w"] >= 600
            f.color.rgb = rgb(r["c"])
            if b["ls"]:
                run._r.get_or_add_rPr().set("spc", str(int(b["ls"] * 75)))   # 1/100 pt


def build_pptx(slides, css, out):
    out = pathlib.Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(W), px(H)
    with tempfile.TemporaryDirectory() as tmp:
        for i, body in enumerate(slides, 1):
            bg = pathlib.Path(tmp) / f"bg{i}.png"
            chrome(page(body, css + NO_TEXT), tmp, f"--screenshot={bg}")
            s = prs.slides.add_slide(prs.slide_layouts[6])
            s.shapes.add_picture(str(bg), 0, 0, px(W), px(H))
            for b in measure(body, css, tmp):
                add_block(s, b)
    prs.save(out)
    return out


def to_canva(path, title):
    from publish.canva import Canva
    cv = Canva()
    return cv.import_file(str(path), title, cv.ensure_folders()["queue"])
