"""핫코노미 카드뉴스(1080×1350, 4:5) PPTX 만들기 → 캔바로 가져가 편집·발행.

    from publish.cardnews import build
    build([{"kind": "cover", "tag": "경제 기초", "title": "...\n**강조**", "sub": "..."},
           {"kind": "body", "title": "...", "body": "...", "point": "핵심 한 줄"},
           {"kind": "end", "title": "정리하면", "checks": ["...", "..."]}], "out.pptx")
    python -m publish.cardnews --sample        # 시안 만들어 캔바 '발행 대기' 폴더로 올리고 JPG 미리보기 저장

디자인(로고에서 따옴): 거의 검정 바탕 + 빨간 불빛, 바닥에 흐린 캔들 차트, 표지엔 올라가는 빨간 화살표.
강조는 불꽃 주황→노랑, 버튼·포인트는 로고 빨강. 배경은 그림(편집 안 함), 글자·도형은 캔바에서 그대로 고칠 수 있다.
글꼴은 캔바에 있는 Noto Sans KR (가져올 때 글꼴이 바뀌지 않게). 로고는 assets/brand/hotconomy_logo.png.
"""
import random
import sys
import urllib.request

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from common import ROOT

PX = 9525                      # 1px (96dpi) = 9525 EMU
W, H = 1080, 1350
FONT = "Noto Sans KR"
RED, ORANGE, YELLOW = RGBColor(0xE3, 0x26, 0x2B), RGBColor(0xFF, 0x7A, 0x1A), RGBColor(0xFF, 0xC2, 0x3A)
WHITE, GREY, DIM, INK = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xB4, 0xB4, 0xBC), RGBColor(0x2A, 0x0E, 0x10), RGBColor(0x0B, 0x0B, 0x0D)
LOGO = ROOT / "assets" / "brand" / "hotconomy_logo.png"
BG_DIR = ROOT / "output" / "cardnews" / "_bg"
DISCLAIMER = "정보 제공용이며 투자 권유가 아닙니다"


def px(v):
    return Emu(int(v * PX))


# ── 배경 그림 (Pillow) ──
def _glow(img, cx, cy, r, color, alpha):
    layer = Image.new("RGBA", img.size, (0, 0, 0, 0))
    ImageDraw.Draw(layer).ellipse((cx - r, cy - r, cx + r, cy + r), fill=(*color, alpha))
    img.alpha_composite(layer.filter(ImageFilter.GaussianBlur(r * 0.45)))


def _candles(img, y_base, height, alpha, seed):
    rnd, d = random.Random(seed), ImageDraw.Draw(img, "RGBA")
    price = y_base - height * 0.3
    for x in range(-10, W + 40, 46):
        move = rnd.uniform(-1, 1.25) * height * 0.12
        o, c = price, price - move
        hi, lo = min(o, c) - rnd.uniform(6, 30), max(o, c) + rnd.uniform(6, 30)
        col = (227, 38, 43, alpha) if c < o else (120, 120, 130, int(alpha * 0.7))
        d.line((x + 13, hi, x + 13, lo), fill=col, width=3)
        d.rectangle((x, min(o, c), x + 26, max(o, c) + 2), fill=col)
        price = max(y_base - height, min(y_base, c))


def background(kind):
    path = BG_DIR / f"{kind}.png"
    if path.exists():
        return path
    BG_DIR.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGBA", (W, H), (11, 11, 13, 255))
    if kind == "cover":
        _glow(img, W - 120, 260, 520, (227, 38, 43), 120)
        _glow(img, 120, H - 160, 380, (255, 90, 20), 60)
        _candles(img, H - 110, 260, 60, 7)
        layer = Image.new("RGBA", img.size, (0, 0, 0, 0))
        d = ImageDraw.Draw(layer)
        pts = [(40, H - 150), (280, H - 250), (450, H - 205), (700, H - 360), (860, H - 320), (1010, H - 470)]
        d.line(pts, fill=(255, 70, 40, 255), width=14, joint="curve")
        d.polygon([(1040, H - 520), (972, H - 482), (1028, H - 430)], fill=(255, 70, 40, 255))
        img.alpha_composite(layer.filter(ImageFilter.GaussianBlur(10)))
        img.alpha_composite(layer)
    elif kind == "body":
        _glow(img, W + 40, -40, 460, (227, 38, 43), 85)
        _candles(img, H - 110, 240, 34, 11)
    else:
        _glow(img, W // 2, 360, 480, (227, 38, 43), 110)
        _candles(img, H - 110, 260, 40, 23)
    img.convert("RGB").save(path)
    return path


# ── 도형·글자 ──
def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb, s.line.width = line, Pt(2.5)
    if radius is not None:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, value, size, color=WHITE, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.3):
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment, p.line_spacing = align, spacing
        hot = line.startswith("**") and line.endswith("**")      # **줄** = 강조색
        r = p.add_run()
        r.text = line.strip("*")
        r.font.name, r.font.size, r.font.bold = FONT, Pt(size * 0.75), bold or hot
        r.font.color.rgb = YELLOW if hot else color
    return tb


def frame(slide, kind, n, total):
    slide.shapes.add_picture(str(background(kind)), 0, 0, px(W), px(H))
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), px(60), px(52), px(84), px(84))
    text(slide, 160, 58, 400, 40, "핫코노미", 34, bold=True)
    text(slide, 160, 100, 400, 30, "@hotconomy", 22, GREY)
    text(slide, W - 260, 70, 200, 50, f"{n:02d} / {total:02d}", 26, ORANGE, bold=True, align=PP_ALIGN.RIGHT)
    text(slide, 60, H - 70, W - 120, 30, DISCLAIMER, 20, GREY, align=PP_ALIGN.CENTER)


def build(slides, out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(W), px(H)
    blank = prs.slide_layouts[6]
    num = 0
    for i, s in enumerate(slides, 1):
        sl = prs.slides.add_slide(blank)
        kind = s["kind"]
        frame(sl, kind, i, len(slides))
        if kind == "cover":
            box(sl, 60, 250, 230, 64, RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
            text(sl, 60, 250, 230, 64, s.get("tag", "경제 기초"), 30, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            text(sl, 60, 350, W - 120, 520, s["title"], 120, bold=True, spacing=1.12)
            text(sl, 60, 880, W - 220, 140, s.get("sub", ""), 44, GREY)
            box(sl, W - 340, 1075, 280, 76, INK, ORANGE, MSO_SHAPE.ROUNDED_RECTANGLE, 0.5)
            text(sl, W - 340, 1075, 280, 76, "넘겨서 확인  ›", 32, ORANGE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        elif kind == "body":
            num += 1
            text(sl, 60, 190, 400, 200, f"{num:02d}", 170, ORANGE, bold=True, spacing=1.0)
            text(sl, 60, 410, W - 120, 200, s["title"], 84, bold=True, spacing=1.15)
            box(sl, 60, 600, 10, 260, RED)
            text(sl, 100, 592, W - 160, 300, s["body"], 56, spacing=1.45)
            if s.get("point"):
                box(sl, 60, 940, W - 120, 200, DIM, RED, MSO_SHAPE.ROUNDED_RECTANGLE, 0.12)
                text(sl, 100, 970, 300, 40, "🔥 핵심", 30, ORANGE, bold=True)
                text(sl, 100, 1020, W - 200, 110, s["point"], 46, bold=True, spacing=1.25)
        else:
            if LOGO.exists():
                sl.shapes.add_picture(str(LOGO), px(W / 2 - 130), px(190), px(260), px(260))
            text(sl, 60, 490, W - 120, 110, s["title"], 80, bold=True, align=PP_ALIGN.CENTER)
            text(sl, 160, 640, W - 320, 320, "\n".join(f"✓  {c}" for c in s.get("checks", [])), 46, spacing=1.6)
            box(sl, 90, 1040, 430, 110, INK, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, 0.5)
            text(sl, 90, 1040, 430, 110, "🔖 저장해두기", 38, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            box(sl, 560, 1040, 430, 110, RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
            text(sl, 560, 1040, 430, 110, "팔로우 하기", 38, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    prs.save(out)
    return out


SAMPLE = [
    {"kind": "cover", "tag": "경제 기초", "title": "금리가 내리면\n내 돈은\n**어떻게 될까?**", "sub": "예금·대출·주식\n딱 3가지만 보면 돼요"},
    {"kind": "body", "title": "예금·적금", "body": "새로 드는 예금 이자는\n낮아져요.\n만기 전에 깨면 손해예요.", "point": "이미 든 고정금리 예금은 그대로!"},
    {"kind": "body", "title": "대출", "body": "변동금리 대출은\n이자 부담이 줄어요.\n고정금리는 따져 보고.", "point": "갈아타기 전 중도상환수수료 확인"},
    {"kind": "body", "title": "주식·채권", "body": "돈 빌리는 비용이 싸지면\n성장주·채권이\n먼저 반응해요.", "point": "'왜 내렸는지'가 더 중요해요"},
    {"kind": "end", "title": "정리하면", "checks": ["예금은 미리 묶어 두기", "대출은 조건부터 확인", "주식은 내린 이유 보기"]},
]


def main():
    from datetime import datetime

    from common import load_env
    from publish.canva import Canva
    load_env()
    out_dir = ROOT / "output" / "cardnews"
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%m%d-%H%M")
    path = build(SAMPLE, out_dir / f"sample-{stamp}.pptx")
    if "--local" in sys.argv:
        print("✅ PPTX 저장:", path)
        return
    cv = Canva()
    folders = cv.ensure_folders()
    design = cv.import_file(str(path), f"[시안] 금리 카드뉴스 {stamp}", folders["queue"])
    print("✅ 캔바 '발행 대기' 폴더에 올림:", design["urls"]["edit_url"])
    for i, url in enumerate(cv.export(design["id"], "jpg"), 1):
        urllib.request.urlretrieve(url, out_dir / f"sample-{stamp}-{i}.jpg")
    print("✅ JPG 미리보기 저장:", out_dir, stamp)


if __name__ == "__main__":
    if "--sample" in sys.argv:
        main()
